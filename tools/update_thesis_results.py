from pathlib import Path
import unicodedata
from docx import Document


SOURCE = Path(r"C:\Users\Usuario\Downloads\Tesis_4ta.docx")
OUTPUT = Path(r"D:\git clone\Sistema Gis\output\Tesis_4ta_actualizada.docx")


REPLACEMENTS = {
    "Respecto al dashboard administrativo,": (
        "Respecto al dashboard administrativo, los resultados demuestran que el sistema ofrece una interfaz funcional "
        "para la gestion de reportes y apoyo a la toma de decisiones. Las pruebas validaron el login administrativo, "
        "la aplicacion de filtros, la tabla de reportes, la paginacion, la exportacion CSV, la apertura del detalle, "
        "las comparaciones temporales y las secciones de analisis estadistico y espacial. Este ultimo punto permite "
        "presentar preguntas comprensibles para gestores y conservar los nombres tecnicos como detalle metodologico. "
        "En el dashboard, los 13 casos automatizados finalizaron correctamente."
    ),
    "Finalmente, los antecedentes internacionales relacionados con analisis espacial": (
        "Finalmente, los antecedentes internacionales relacionados con analisis espacial respaldan la utilidad de "
        "Moran's I, Getis-Ord Gi*, mapas de calor, hotspots y coldspots para identificar patrones territoriales. En "
        "la implementacion se generaron cuadriculas de 250 y 500 metros con vecindad Queen. Sobre 64 unidades de "
        "500 metros, Moran's I global obtuvo I=0.309265 y p=0.005 con 999 permutaciones, evidenciando autocorrelacion "
        "espacial positiva. Moran local y Getis-Ord Gi* no conservaron unidades significativas despues de aplicar "
        "Benjamini-Hochberg FDR 0.05, resultado coherente con el reducido volumen actual. Estas tecnicas se presentan "
        "como analisis espacial exploratorio y no como prediccion delictiva."
    ),
    "No obstante, los resultados deben interpretarse considerando las limitaciones.": (
        "No obstante, los resultados deben interpretarse considerando las limitaciones. Las pruebas se ejecutaron en "
        "entornos controlados y no sustituyen una validacion con ciudadanos y gestores reales. El volumen actual de "
        "reportes limita la potencia de Moran local y Getis-Ord Gi*. Los sectores A-G fueron digitalizados de forma "
        "referencial desde el mapa raster municipal de 2019 y requieren cartografia vectorial oficial para uso "
        "catastral. Ademas, el almacenamiento local de fotografias requiere politicas de respaldo, retencion y "
        "seguridad antes de un despliegue institucional."
    ),
    "\u25cf Segundo:": (
        "\u25cf Segundo: Se logro construir la infraestructura GIS con PostgreSQL/PostGIS, sectores territoriales A-G, "
        "asignacion espacial automatica, cuadriculas de 250 y 500 metros y matrices de vecindad Queen. El dashboard "
        "React con Leaflet y OpenStreetMap permite visualizar capas configurables, mapas de calor, concentraciones "
        "descriptivas, Moran's I global y local, y Getis-Ord Gi*, manteniendo advertencias sobre significancia, "
        "causalidad y disponibilidad de datos."
    ),
    "\u25cf Tercero:": (
        "\u25cf Tercero: Se logro desarrollar el modulo de analisis y consulta, con filtros por categoria, fecha, "
        "estado, urgencia y zona, indicadores visuales, comparaciones entre periodos y exportacion CSV. El dashboard "
        "incorpora analisis estadistico no parametrico y categorico, asi como analisis espacial exploratorio, "
        "presentados mediante preguntas comprensibles para gestores y sin plantear el sistema como una herramienta "
        "de prediccion delictiva."
    ),
    "A partir de los resultados obtenidos, se identifican como trabajos futuros": (
        "A partir de los resultados obtenidos, se identifican como trabajos futuros la validacion del sistema en "
        "condiciones reales de uso, la evaluacion de utilidad con gestores municipales, la medicion de tiempos reales "
        "de reporte y atencion, la obtencion de cartografia vectorial oficial para los sectores A-G, el incremento del "
        "volumen de observaciones para fortalecer Moran local y Getis-Ord Gi*, y la integracion progresiva con "
        "procesos institucionales de seguridad ciudadana. Estas ampliaciones deben mantener el enfoque de analisis "
        "territorial exploratorio y apoyo a decisiones, sin atribuir causalidad ni orientar el sistema a modelos "
        "predictivos."
    ),
}


def replace_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def comparable(text):
    normalized = unicodedata.normalize("NFKD", text.replace("\u2019", "'"))
    return "".join(character for character in normalized if not unicodedata.combining(character))


def main():
    document = Document(SOURCE)
    pending = dict(REPLACEMENTS)

    for paragraph in document.paragraphs:
        normalized = comparable(paragraph.text)
        for prefix, replacement in list(pending.items()):
            normalized_prefix = comparable(prefix)
            if normalized.startswith(normalized_prefix):
                replace_text(paragraph, replacement)
                del pending[prefix]
                break

    if pending:
        raise RuntimeError(f"No se encontraron los parrafos: {list(pending)}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Updated {len(REPLACEMENTS)} paragraphs in {OUTPUT}")


if __name__ == "__main__":
    main()
